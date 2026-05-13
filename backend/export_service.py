import io
from typing import Dict, List
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def export_to_word(filename: str, requirements: List[dict], selections: Dict[str, str]) -> bytes:
    doc = Document()

    doc.add_heading("Cahier des Charges Reformulé", 0)
    doc.add_paragraph(f"Document original: {filename}").style = "Normal"
    doc.add_paragraph("")

    for req in requirements:
        req_id = req.get("id", "?")
        nom = req.get("nom", "Sans titre")
        selected = selections.get(req_id, req.get("contenu", ""))

        doc.add_heading(f"{req_id} — {nom}", level=2)
        doc.add_paragraph(selected).style = "Normal"
        doc.add_paragraph("")

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()


def export_to_excel(filename: str, requirements: List[dict], reformulations: Dict = None, selections: Dict[str, str] = None) -> bytes:
    selections = selections or {}
    reformulations = reformulations or {}

    wb = Workbook()
    ws = wb.active
    ws.title = "Exigences"

    header_fill = PatternFill(start_color="E8A020", end_color="E8A020", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    headers = ["ID", "Nom", "Contenu Original", "Conforme?", "Reformulation Retenue", "Option 1", "Option 2", "Option 3"]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 40
    ws.column_dimensions["F"].width = 30
    ws.column_dimensions["G"].width = 30
    ws.column_dimensions["H"].width = 30

    for req in requirements:
        req_id = req.get("id", "?")
        nom = req.get("nom", "")
        contenu = req.get("contenu", "")
        selected = selections.get(req_id, contenu)

        refo_data = reformulations.get(req_id, {})
        conforme = "Oui" if refo_data.get("conforme") else "Non"
        options = refo_data.get("reformulations", [])

        row = [
            req_id,
            nom,
            contenu,
            conforme,
            selected,
            options[0] if len(options) > 0 else "",
            options[1] if len(options) > 1 else "",
            options[2] if len(options) > 2 else "",
        ]

        ws.append(row)

        for cell in ws[ws.max_row]:
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
